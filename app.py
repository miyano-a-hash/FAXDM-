# -*- coding: utf-8 -*-
"""
FAXDM 校正ツール（やどかり） ― Streamlit アプリ

PDF原稿をアップロードすると、月別マスタ（YYMM_やどかり）と FaxDM訴求文 を参照して
次の項目をチェックし、結果を画面とExcelで返します。

  ① 日付と曜日が合っているか
  ② メニュー名がマスタと完全一致しているか（メイン/サブ/副菜2種/サラダ、贅沢あり版は贅沢弁当）
  ③ 日付の上のアイコン（プレミアム/スペシャル/フェア/贅沢プレミアム）と名前がマスタと合っているか
  ④ イチオシ欄の日付・曜日・メニュー・一言コメントが FaxDM訴求文 と合っているか
  ⑤ 贅沢あり版：「●上記メニューに加え、副菜2種とサラダ（もしくはデザート）がつきます。」があるか
  ⑥ 贅沢あり版：混ぜご飯・混ぜ麺・ふりかけの日にマークと注記があるか
  ⑦ 掲載期間（ファイル名）と掲載日の整合、掲載漏れ
  ⑧ 共通の定型文言（「お得情報も満載」など）

起動:  streamlit run app.py
"""
from __future__ import annotations

import datetime as dt
import io
import re
import statistics
import unicodedata
from dataclasses import dataclass, field
from typing import Optional

import pandas as pd
import pymupdf  # PyMuPDF
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ============================================================================
# 設定（必要に応じてここを書き換えてください）
# ============================================================================
# Googleドライブ：月別マスタが入っているフォルダと FaxDM訴求文 のID
# （Streamlit の Secrets に drive_folder_id / faxdm_sheet_id があればそちらを優先）
DRIVE_FOLDER_ID = "1NVYeVYTyNCW6fLL76pum7MU9cxjvAs-b"
FAXDM_SHEET_ID = "1jqFSHwmtZRYmTlx76Bm2SlTPwAoI8MNctiFNXudW13o"

# 月別マスタのスプレッドシート（名前 → ID）。新しい月ができたら1行追加してください。
# ID は URL の https://docs.google.com/spreadsheets/d/【ここ】/edit の部分です。
# ※ どのシートも共有設定を「リンクを知っている全員」が「閲覧者」にしておく必要があります。
MASTER_SHEETS = {
    "2601_やどかり": "1f6YNPxMnES1N1xAjt5ZZR5JmpUSIcGkR5lT6MZcAjNc",
    "2602_やどかり": "1P_6EQ_sEUWBxt5T7uyXx_tIzqjSYLOoYoGtHfzadqIY",
    "2603_やどかり": "1Ek0E9p11eMmHxhq0s_VsrcW3MjLV3aWbWvglHj8_Elo",
    "2604_やどかり": "1eTdNQUTIM-mQOqAh5hjRdcLIorhPH8r5xGfxuOqbEqs",
    "2605_やどかり": "1FQjoeok9M7aR82ybyaCEZLO4x1x16H4pov7YP9e_NnY",
    "2606_やどかり": "1UvB4Ix7Ewj32RUo37nohhmozqPzzmXzAXoAO6bvQ1Bw",
    "2607_やどかり": "1ZPRkSi2wo0LSO7v7rmo32e2Yao4bEj9ofcHFeY4W8cE",
    "2608_やどかり": "1KSxFK-1PN_7JSD-FL8SKOpJiZ2d_UpnKB99MGArjWbo",
    "2609_やどかり": "1GQnzGk5H9mETyDJwrZoxs6sMPZa1F54UTdy5rJcK4lM",
    "2610_やどかり": "1aOGW5eG5N43JFl9xcvqRryrloPv_OnCn3vXn67Sx8BY",
    "2611_やどかり": "18eV1VJ6rQtB8UQ_oFzBNVCd2DUi6Pg5ajRvmC0aGpDw",
    "2612_やどかり": "14NRgWd1mQ8l_zKQ_XHTQyUvhNQf3h4v3CIRqLN9YWGI",
}

# 月別マスタのファイル名（{yy}=年下2桁, {mm}=月2桁）。例: 2610_やどかり
MONTHLY_NAME = "{yy}{mm}_やどかり"
FAXDM_NAME = "FaxDM訴求文"

# マスタのラベル列（日付の左隣の列）の塗り色 → アイコン種別（近い色＝色差24以内も同じ扱い）
ICON_COLORS = {
    "FFFF00": "プレミアム",      # 黄色
    "EAD1DC": "贅沢プレミアム",  # うすいピンク
    "B4C6E7": "スペシャル",      # 薄い青
    "FBBC04": "フェア",          # オレンジ
}
# ラベルにこの語が含まれる日を「ご飯変更不可の日」とする
RICE_WORDS = ["混ぜご飯", "混ぜ麺", "ふりかけ"]
# 贅沢あり版に必要な文言
ZEITAKU_REQUIRED = ["●上記メニューに加え、副菜2種とサラダ（もしくはデザート）がつきます。"]
# 全版共通で必要な文言
COMMON_REQUIRED = ["お得情報も満載", "あなたの街のお弁当屋さん"]

WEEKDAY_CHARS = {"㊊": 0, "㊋": 1, "㊌": 2, "㊍": 3, "㊎": 4, "㊏": 5, "㊐": 6}
WEEKDAY_JA = "月火水木金土日"
TYPE_WORDS = ["贅沢プレミアム", "プレミアム", "スペシャル", "フェア"]
TYPE_ALIASES = {"SP": "スペシャル", "ＳＰ": "スペシャル", "贅沢プレ": "贅沢プレミアム",
                "スペシャル": "スペシャル", "プレミアム": "プレミアム", "フェア": "フェア",
                "贅沢プレミアム": "贅沢プレミアム"}
# アイコンの文字がアウトライン化されている場合の判定用（文字ごとのパス数の並び）
ICON_GLYPH_SIGNATURES = {
    (17, 7, 12, 12, 16): "プレミアム",
    (11, 19, 13, 14, 16): "スペシャル",
    (8, 12, 12): "フェア",
}
MENU_FIELDS = ["main", "submain", "fukusai1", "fukusai2", "salad"]
ZEITAKU_FIELDS = ["zeitaku1", "zeitaku2"]
FIELD_JA = {"main": "メイン", "submain": "サブメイン", "fukusai1": "副菜1", "fukusai2": "副菜2",
            "salad": "サラダ(デザート)", "zeitaku1": "贅沢弁当1", "zeitaku2": "贅沢弁当2"}

SEV_ORDER = {"高": 0, "中": 1, "低": 2, "情報": 3}


# ----------------------------------------------------------------------------
# 文字列ユーティリティ
# ----------------------------------------------------------------------------
def fix_extract(s: str) -> str:
    """PDF抽出時に入る「数字の後ろの空白」（例: 1 品, 4 種）を取り除く。"""
    s = re.sub(r"(?<=[0-9]) (?=[^\x00-\x7f])", "", s)
    return s.strip()


def no_ws(s: str) -> str:
    return re.sub(r"[\s　]+", "", s or "")


def loose(s: str) -> str:
    """スペース・波ダッシュの字形差を無視した比較用文字列。"""
    s = no_ws(s)
    return s.replace("～", "〜")


def describe_ws(s: str) -> str:
    return s.replace("　", "［全角空白］").replace(" ", "［半角空白］") if s else s


def compare_text(pdf: str, master: str, wrapped: bool = False):
    """戻り値: (判定, 説明)  判定は 'OK' / '低' / '高'"""
    p = fix_extract(pdf or "")
    m = (master or "").strip()
    if p == m:
        return "OK", ""
    if no_ws(p) == no_ws(m):
        if wrapped:
            return "OK", ""
        return "低", "スペースの種類・有無だけの差（原稿を目視確認）"
    if loose(p) == loose(m):
        return "低", "「～」と「〜」など記号の字形差"
    if unicodedata.normalize("NFKC", no_ws(p)) == unicodedata.normalize("NFKC", no_ws(m)):
        return "高", "全角・半角の違い（括弧「（）」と「()」など）"
    return "高", ""



def split_sides(lines: list, master_vals: list):
    """副菜・サラダ行を「/」で3つに分ける。
    行末に「/」がなく改行で区切られている場合も試し、マスタに合う分け方を採用する。
    戻り値: (items[(text, wrapped)], 改行を区切りとして扱ったか)"""
    import itertools
    breaks = []  # 「/」に隣接しない改行の位置
    for i in range(len(lines) - 1):
        a, b = lines[i].rstrip(), lines[i + 1].lstrip()
        if not a.endswith("/") and not b.startswith("/"):
            breaks.append(i)

    def build(seps):
        buf = lines[0]
        for i in range(len(lines) - 1):
            buf += ("/" if i in seps else "\n") + lines[i + 1]
        parts = [p for p in buf.split("/")]
        return [(p.replace("\n", "").strip(), "\n" in p.strip()) for p in parts if p.strip()]

    base = build(set())
    if len(base) == len(master_vals) or not breaks:
        return base, False
    best = None
    for k in range(1, len(breaks) + 1):
        for comb in itertools.combinations(breaks, k):
            items = build(set(comb))
            if len(items) == len(master_vals):
                score = sum(no_ws(a[0]) == no_ws(b) for a, b in zip(items, master_vals))
                if best is None or score > best[0]:
                    best = (score, items)
    if best and best[0] > 0:
        return best[1], True
    return base, False

# ----------------------------------------------------------------------------
# 設定オブジェクト
# ----------------------------------------------------------------------------
@dataclass
class Config:
    icon_colors: dict = field(default_factory=lambda: dict(ICON_COLORS))
    rice_words: list = field(default_factory=lambda: list(RICE_WORDS))
    zeitaku_required: list = field(default_factory=lambda: list(ZEITAKU_REQUIRED))
    common_required: list = field(default_factory=lambda: list(COMMON_REQUIRED))
    monthly_name: str = MONTHLY_NAME
    faxdm_name: str = FAXDM_NAME


# ----------------------------------------------------------------------------
# マスタ
# ----------------------------------------------------------------------------
@dataclass
class MasterDay:
    date: dt.date
    fields: dict
    label: str
    label_color: Optional[str]
    icon_type: Optional[str]     # プレミアム/スペシャル/フェア/贅沢プレミアム/None
    event_name: str               # アイコンの名前（ハロウィン など）
    rice_word: Optional[str]      # 混ぜご飯/ふりかけ/混ぜ麺
    memo: str                     # ラベル列より左のメモ
    source: str

    @property
    def is_full(self) -> bool:
        return bool((self.fields.get("submain") or "").strip())


@dataclass
class FaxRow:
    month: int
    day: int
    type: str
    event: str
    menu: str
    fixed: str
    adopt: str


def _rgb_of(cell) -> Optional[str]:
    f = cell.fill
    if f is None or f.fill_type is None:
        return None
    c = f.fgColor
    if c is not None and c.type == "rgb" and isinstance(c.rgb, str):
        return c.rgb[-6:].upper()
    return None


def _color_to_type(rgb: Optional[str], table: dict) -> Optional[str]:
    if not rgb:
        return None
    if rgb in table:
        return table[rgb]
    r, g, b = int(rgb[0:2], 16), int(rgb[2:4], 16), int(rgb[4:6], 16)
    best, bd = None, 1e9
    for k, v in table.items():
        try:
            r2, g2, b2 = int(k[0:2], 16), int(k[2:4], 16), int(k[4:6], 16)
        except ValueError:
            continue
        d = ((r - r2) ** 2 + (g - g2) ** 2 + (b - b2) ** 2) ** 0.5
        if d < bd:
            best, bd = v, d
    return best if bd <= 24 else None


def parse_monthly_master(src, source_name: str, year: int, month: int, cfg: Config) -> dict:
    wb = load_workbook(src)
    ws = wb.worksheets[0]  # 一番左のシートだけを見る
    header_row, cols = None, {}
    for r in range(1, min(ws.max_row, 15) + 1):
        vals = {str(ws.cell(r, c).value).strip().lower(): c for c in range(1, min(ws.max_column, 40) + 1)
                if ws.cell(r, c).value is not None}
        if "day" in vals:
            header_row, cols = r, vals
            break
    if header_row is None:
        raise ValueError(f"{source_name}: 一番左のシートに 'day' 見出しが見つかりません")
    day_col = cols["day"]
    label_col = day_col - 1 if day_col > 1 else None
    out = {}
    for r in range(header_row + 1, ws.max_row + 1):
        v = ws.cell(r, day_col).value
        if v is None:
            continue
        if isinstance(v, dt.datetime):
            d = v.date()
        elif isinstance(v, dt.date):
            d = v
        elif isinstance(v, (int, float)) and 1 <= int(v) <= 31:
            try:
                d = dt.date(year, month, int(v))
            except ValueError:
                continue
        else:
            continue
        flds = {}
        for f in MENU_FIELDS + ZEITAKU_FIELDS:
            c = cols.get(f)
            val = ws.cell(r, c).value if c else None
            flds[f] = "" if val is None else str(val)
        label, color, itype, event, rice = "", None, None, "", None
        if label_col:
            lc = ws.cell(r, label_col)
            label = "" if lc.value is None else str(lc.value).strip()
            color = _rgb_of(lc)
            itype = _color_to_type(color, cfg.icon_colors)
            for w in cfg.rice_words:
                if w in label:
                    rice = w
                    break
            if rice and itype:
                itype = None
            if itype in ("プレミアム", "フェア"):
                event = label
            elif itype and label not in TYPE_ALIASES:
                event = label
        memo = " / ".join(str(ws.cell(r, c).value).strip() for c in range(1, (label_col or day_col))
                          if ws.cell(r, c).value is not None and not isinstance(ws.cell(r, c).value, (int, float)))
        out[d] = MasterDay(d, flds, label, color, itype, event, rice, memo, source_name)
    return out


def parse_faxdm(src) -> list:
    wb = load_workbook(src, data_only=True)
    rows = []
    for ws in wb.worksheets:
        hdr, cols = None, {}
        for r in range(1, min(ws.max_row, 10) + 1):
            vals = {str(ws.cell(r, c).value).strip(): c for c in range(1, min(ws.max_column, 30) + 1)
                    if ws.cell(r, c).value is not None}
            if "月" in vals and "日程" in vals:
                hdr, cols = r, vals
                break
        if hdr is None:
            continue

        def g(r, name):
            c = cols.get(name)
            v = ws.cell(r, c).value if c else None
            return "" if v is None else str(v).strip()

        for r in range(hdr + 1, ws.max_row + 1):
            mm = re.search(r"(\d+)", g(r, "月"))
            dd = re.search(r"(\d+)", g(r, "日程"))
            if not (mm and dd):
                continue
            t = g(r, "種別")
            rows.append(FaxRow(int(mm.group(1)), int(dd.group(1)), TYPE_ALIASES.get(t, t), g(r, "イベント名"),
                               g(r, "メニュー名"), g(r, "確定"), g(r, "選定可否")))
        break
    return rows


class MasterSource:
    """マスタの取得元。provider は name -> (bytes, ファイル名) を返す関数。"""

    def __init__(self, provider, cfg: Config):
        self.provider = provider
        self.cfg = cfg
        self.notes = []

    def monthly(self, year: int, month: int) -> dict:
        name = self.cfg.monthly_name.format(yy=f"{year % 100:02d}", mm=f"{month:02d}")
        try:
            got = self.provider(name)
        except Exception as e:
            self.notes.append(f"月別マスタ「{name}」を読み込めませんでした：{e}")
            return {}
        if not got:
            self.notes.append(f"月別マスタ「{name}」が見つかりません")
            return {}
        data, fname = got
        return parse_monthly_master(io.BytesIO(data), fname, year, month, self.cfg)

    def faxdm(self) -> list:
        try:
            got = self.provider(self.cfg.faxdm_name)
        except Exception as e:
            self.notes.append(f"「{self.cfg.faxdm_name}」を読み込めませんでした：{e}（イチオシの訴求文チェックは省略）")
            return []
        if not got:
            self.notes.append(f"「{self.cfg.faxdm_name}」が見つかりません（イチオシの訴求文チェックは省略）")
            return []
        return parse_faxdm(io.BytesIO(got[0]))


def upload_provider(files: dict):
    """アップロードされた .xlsx（ファイル名 -> bytes）から探す"""
    norm = {unicodedata.normalize("NFC", k): v for k, v in files.items()}

    def get(name: str):
        for fname in sorted(norm):
            if fname.startswith(name) and fname.lower().endswith(".xlsx"):
                return norm[fname], fname
        return None
    return get


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def sheet_id_from(text: str) -> Optional[str]:
    """スプレッドシートのURL（またはID）からIDを取り出す"""
    text = (text or "").strip()
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9_-]{20,})", text)
    if m:
        return m.group(1)
    return text if re.fullmatch(r"[A-Za-z0-9_-]{20,}", text) else None


def _guess_sheet_name(data: bytes) -> str:
    """ファイル名が取れなかったときに、中身から「2610_やどかり」「FaxDM訴求文」を推定する"""
    wb = load_workbook(io.BytesIO(data), read_only=True)
    for ws in wb.worksheets[:3]:
        m = re.search(r"(\d{2})年\s*(\d{1,2})月号", ws.title)
        if m:
            return MONTHLY_NAME.format(yy=m.group(1), mm=f"{int(m.group(2)):02d}")
    for row in wb.worksheets[0].iter_rows(min_row=1, max_row=5, values_only=True):
        vals = {str(v).strip() for v in row if v is not None}
        if "月" in vals and "日程" in vals:
            return FAXDM_NAME
    return ""


@st.cache_data(ttl=300, show_spinner=False)
def fetch_sheet(sheet_id: str):
    """共有リンクが有効なスプレッドシートを .xlsx として取得する → (bytes, 名前)"""
    import urllib.parse
    import requests
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    r = requests.get(url, timeout=60)
    if r.status_code != 200 or not r.content.startswith(b"PK"):
        raise RuntimeError("取得できませんでした。共有設定が「リンクを知っている全員（閲覧者）」になっているか確認してください")
    cd = r.headers.get("Content-Disposition", "")
    name = ""
    m = re.search(r"filename\*=UTF-8''([^;]+)", cd)
    if m:
        name = urllib.parse.unquote(m.group(1))
    else:
        m = re.search(r'filename="([^"]+)"', cd)
        if m:
            name = m.group(1)
    name = unicodedata.normalize("NFC", re.sub(r"\.xlsx$", "", name.strip()))
    if not re.match(r"^\d{4}_|^" + re.escape(FAXDM_NAME), name):
        name = _guess_sheet_name(r.content) or name
    return r.content, name


def sheet_provider(master_sheets: dict, extra_urls: list, faxdm_id: str, faxdm_name: str):
    """スプレッドシートのURLから読む。月別マスタは名前で探し、登録がなければ追加URLの中から探す"""
    extra_ids = [i for i in (sheet_id_from(u) for u in extra_urls) if i]

    def get(name: str):
        if name == faxdm_name:
            fid = sheet_id_from(faxdm_id)
            if not fid:
                return None
            data, _ = fetch_sheet(fid)
            return data, faxdm_name
        for sid in extra_ids:  # 追加URLを優先（登録済みの月を差し替えたいとき用）
            data, n = fetch_sheet(sid)
            if n.startswith(name):
                return data, n
        sid = sheet_id_from(master_sheets.get(name, ""))
        if sid:
            data, _ = fetch_sheet(sid)
            return data, name
        return None
    return get


def _drive_service(sa_info: dict):
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    creds = service_account.Credentials.from_service_account_info(
        sa_info, scopes=["https://www.googleapis.com/auth/drive.readonly"])
    return build("drive", "v3", credentials=creds, cache_discovery=False)


@st.cache_data(ttl=600, show_spinner=False)
def _drive_list(sa_json: str, folder_id: str) -> dict:
    """フォルダ（と直下のサブフォルダ）にあるスプレッドシート: 名前 -> (id, 更新日時)"""
    import json
    svc = _drive_service(json.loads(sa_json))
    out, folders = {}, [folder_id]
    opts = dict(supportsAllDrives=True, includeItemsFromAllDrives=True, pageSize=1000)
    sub = svc.files().list(q=f"'{folder_id}' in parents and mimeType='application/vnd.google-apps.folder' "
                             "and trashed=false", fields="files(id,name)", **opts).execute().get("files", [])
    folders += [f["id"] for f in sub]
    for fid in folders:
        res = svc.files().list(q=f"'{fid}' in parents and mimeType='application/vnd.google-apps.spreadsheet' "
                                 "and trashed=false", fields="files(id,name,modifiedTime)", **opts).execute()
        for f in res.get("files", []):
            out.setdefault(unicodedata.normalize("NFC", f["name"]), (f["id"], f["modifiedTime"]))
    return out


@st.cache_data(ttl=3600, show_spinner=False)
def _drive_export(sa_json: str, file_id: str, modified: str) -> bytes:
    import json
    svc = _drive_service(json.loads(sa_json))
    return svc.files().export(fileId=file_id, mimeType=XLSX_MIME).execute()


def drive_provider(sa_info: dict, folder_id: str, faxdm_id: str, faxdm_name: str):
    import json
    sa_json = json.dumps(dict(sa_info), sort_keys=True)

    def get(name: str):
        if name == faxdm_name and faxdm_id:
            return _drive_export(sa_json, faxdm_id, dt.datetime.now().strftime("%Y%m%d%H%M")[:11]), name
        files = _drive_list(sa_json, folder_id)
        hit = files.get(name)
        if not hit:
            return None
        return _drive_export(sa_json, hit[0], hit[1]), name
    return get


# ----------------------------------------------------------------------------
# PDF 読み取り
# ----------------------------------------------------------------------------
@dataclass
class Span:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float
    size: float
    font: str
    color: int

    @property
    def yc(self):
        return (self.y0 + self.y1) / 2

    @property
    def xc(self):
        return (self.x0 + self.x1) / 2

    @property
    def white(self):
        return self.color == 0xFFFFFF

    @property
    def regular(self):
        return "Regu" in self.font or "Light" in self.font or "-L" in self.font

    @property
    def din(self):
        return "DIN" in self.font


@dataclass
class PdfDay:
    day: int
    month: Optional[int] = None
    year: Optional[int] = None
    weekday_char: str = ""
    col: int = 0
    yc: float = 0.0
    top: float = 0.0
    items: list = field(default_factory=list)       # [(text, wrapped, rice_flag)]
    sides: list = field(default_factory=list)
    sides_raw: str = ""
    side_lines: list = field(default_factory=list)
    zeitaku: list = field(default_factory=list)     # [(text, wrapped)]
    zeitaku_has_slash: bool = False
    icon_type: Optional[str] = None
    icon_how: str = ""
    icon_name: str = ""
    rice_marks: list = field(default_factory=list)

    @property
    def date(self) -> Optional[dt.date]:
        try:
            return dt.date(self.year, self.month, self.day)
        except (TypeError, ValueError):
            return None


@dataclass
class Ichioshi:
    day: int
    month: Optional[int]
    weekday_char: str
    type: str
    event: str
    menu: str
    menu_wrapped: bool
    comment: str
    year: Optional[int] = None

    @property
    def date(self):
        try:
            return dt.date(self.year, self.month, self.day)
        except (TypeError, ValueError):
            return None


@dataclass
class PdfDoc:
    path: str
    title: str
    zeitaku: bool
    period: Optional[tuple]
    days: list
    ichioshi: list
    fulltext: str
    footnote_text: str
    warnings: list


def _spans(page) -> list:
    out = []
    for b in page.get_text("dict")["blocks"]:
        for l in b.get("lines", []):
            for s in l["spans"]:
                t = s["text"]
                if not t.strip():
                    continue
                x0, y0, x1, y1 = s["bbox"]
                out.append(Span(t, x0, y0, x1, y1, s["size"], s["font"], s["color"]))
    return out


def _period_from(text: str):
    """'261023配信261027-1106' → (配信日, 開始日, 終了日)"""
    t = unicodedata.normalize("NFKC", text)
    m = re.search(r"(\d{2})(\d{2})(\d{2})\s*配信\s*(\d{2})(\d{2})(\d{2})\s*-\s*(\d{2})(\d{2})", t)
    if not m:
        return None
    y1, m1, d1, y2, m2, d2, m3, d3 = map(int, m.groups())
    try:
        start = dt.date(2000 + y2, m2, d2)
        end_y = start.year + (1 if m3 < m2 else 0)
        return dt.date(2000 + y1, m1, d1), start, dt.date(end_y, m3, d3)
    except ValueError:
        return None


def _group_lines(spans: list, tol: float = 3.0) -> list:
    lines = []
    for s in sorted(spans, key=lambda s: (s.yc, s.x0)):
        for ln in lines:
            if abs(ln["yc"] - s.yc) <= tol:
                ln["spans"].append(s)
                break
        else:
            lines.append({"yc": s.yc, "spans": [s]})
    for ln in lines:
        ln["spans"].sort(key=lambda s: s.x0)
    lines.sort(key=lambda l: l["yc"])
    return lines


def _join(spans: list) -> str:
    out = ""
    prev = None
    for s in spans:
        if prev is not None and s.x0 - prev.x1 > max(prev.size, s.size) * 0.6 and not out.endswith((" ", "　", "/")):
            out += " "
        out += s.text
        prev = s
    return out


def _icon_from_glyphs(glyphs: list) -> Optional[str]:
    sig = tuple(glyphs)
    if sig in ICON_GLYPH_SIGNATURES:
        return ICON_GLYPH_SIGNATURES[sig]
    if len(sig) == 7 and sig[2:] == (17, 7, 12, 12, 16):
        return "贅沢プレミアム"
    best, bd = None, 1e9
    for k, v in ICON_GLYPH_SIGNATURES.items():
        if len(k) == len(sig):
            d = sum(abs(a - b) for a, b in zip(k, sig))
            if d < bd:
                best, bd = v, d
    if best and bd <= 6:
        return best
    if len(sig) == 7:
        return "贅沢プレミアム"
    return None


def read_pdf(data: bytes, name: str, cfg: Config) -> PdfDoc:
    doc = pymupdf.open(stream=data, filetype="pdf")
    page = doc[0]
    title = (doc.metadata or {}).get("title") or ""
    spans = _spans(page)
    warnings = []
    if len(doc) > 1:
        warnings.append(f"{len(doc)}ページありますが、1ページ目だけをチェックしました")
    if not any(re.search(r"[ぁ-んァ-ン一-龥]", s.text) for s in spans):
        raise ValueError("PDFから日本語の文字を読み取れませんでした（文字がアウトライン化されている可能性があります）")

    fulltext = "".join(fix_extract(s.text) for s in sorted(spans, key=lambda s: (round(s.y0), s.x0)))
    period = _period_from(name) or _period_from(title)

    # --- グリッド（注文欄）の位置 ---
    headers = sorted([s for s in spans if re.match(r"^\s*\d{4}\s*年", s.text)], key=lambda s: s.x0)
    if not headers:
        raise ValueError("注文欄の見出し（2026 年◯月）が見つかりません")
    grid_top = min(h.y0 for h in headers)
    col_lefts = [h.x0 for h in headers]
    foot = [s for s in spans if s.y0 > grid_top and ("※衛生面" in s.text or s.text.startswith("●上記メニュー"))]
    grid_bottom = min((s.y0 for s in foot), default=page.rect.height)
    page_w = page.rect.width

    def col_of(x: float) -> int:
        idx = 0
        for i, cl in enumerate(col_lefts):
            if x >= cl - 6:
                idx = i
        return idx

    def col_right(i: int) -> float:
        return col_lefts[i + 1] - 3 if i + 1 < len(col_lefts) else page_w

    zeitaku_edition = any(s.text.strip() == "贅沢弁当" and s.y0 > grid_top - 5 for s in spans)
    weekday_spans = [s for s in spans if s.text.strip() in WEEKDAY_CHARS]
    din_digits = [s for s in spans if s.din and s.text.strip().isdigit()]

    # --- グリッドの日付 ---
    days = []
    used = set()
    for s in din_digits:
        if not (grid_top < s.y0 < grid_bottom) or s.size < 15:
            continue
        wd = [w for w in weekday_spans if s.y0 + s.size * 0.6 <= w.y0 <= s.y1 + 10
              and w.x0 < s.x1 + 3 and w.x1 > s.x0 - 3]
        if not wd:
            continue
        w = min(wd, key=lambda w: w.y0)
        pd_ = PdfDay(day=int(s.text.strip()), weekday_char=w.text.strip(), col=col_of(s.x0), yc=s.yc, top=s.y0)
        pref = [p for p in din_digits if p is not s and p.size < s.size * 0.75 and s.x0 - 25 < p.x1 <= s.x0 + 1
                and abs(p.yc - s.yc) < s.size]
        if pref:
            pd_.month = int(pref[0].text.strip())
            used.add(id(pref[0]))
        used.update({id(s), id(w)})
        days.append(pd_)
    if not days:
        raise ValueError("注文欄の日付を読み取れませんでした")

    # 年月の割り当て（見出しの月から順に、日が戻ったら翌月）
    hdr0 = unicodedata.normalize("NFKC", headers[0].text)
    year = int(re.search(r"(\d{4})", hdr0).group(1))
    months = [int(x) for x in re.findall(r"(\d{1,2})\s*月", hdr0)] or \
             [int(x) for x in re.findall(r"(?<=[年/・])\s*(\d{1,2})", hdr0)]
    m_nums = re.findall(r"年\s*([\d\s/・]+)月", hdr0)
    if m_nums:
        months = [int(x) for x in re.findall(r"\d+", m_nums[0])]
    cur_m = months[0] if months else (period[1].month if period else 1)
    cur_y = year
    prev_day = None
    days.sort(key=lambda d: (d.col, d.yc))
    for d in days:
        if d.month:
            if d.month < cur_m:
                cur_y += 1
            cur_m = d.month
        elif prev_day is not None and d.day < prev_day:
            cur_m += 1
            if cur_m > 12:
                cur_m, cur_y = 1, cur_y + 1
        d.month, d.year = cur_m, cur_y
        prev_day = d.day

    # --- セル範囲 ---
    tabs = []
    for dr in page.get_drawings():
        r = dr["rect"]
        fill = dr.get("fill")
        if fill and sum(fill) < 0.6 and 24 <= r.width <= 70 and 12 <= r.height <= 20 and grid_top < r.y0 < grid_bottom:
            tabs.append({"rect": r, "glyphs": []})
    for dr in page.get_drawings():
        r = dr["rect"]
        fill = dr.get("fill")
        if fill and min(fill) > 0.95 and r.width < 12 and r.height < 12:
            for t in tabs:
                if t["rect"].contains(r):
                    t["glyphs"].append((r.x0, len(dr["items"])))
    for t in tabs:
        t["glyphs"] = [n for _, n in sorted(t["glyphs"])]

    # 注文欄の横罫線（セルの上端）を集める
    hlines = []
    for dr in page.get_drawings():
        for it in dr["items"]:
            if it[0] == "l":
                a, b = it[1], it[2]
                if abs(a.y - b.y) < 0.8 and abs(a.x - b.x) > 60:
                    hlines.append((a.y, min(a.x, b.x), max(a.x, b.x)))
            elif it[0] == "re":
                r = it[1]
                fill = dr.get("fill")
                if r.width > 60 and (fill is None or r.height < 1.5 or min(fill) > 0.9):
                    hlines.append((r.y0, r.x0, r.x1))
                    hlines.append((r.y1, r.x0, r.x1))

    cell_bounds = {}
    for c in range(len(col_lefts)):
        cd = sorted([d for d in days if d.col == c], key=lambda d: d.yc)
        if not cd:
            continue
        cl = col_lefts[c]
        cw = col_right(c) - cl
        ys = sorted(y for y, x0, x1 in hlines
                    if grid_top - 2 < y < grid_bottom + 2 and x0 <= cl + 12 and x1 >= cl + cw * 0.5)
        uniq = []
        for y in ys:
            if not uniq or y - uniq[-1] > 1.5:
                uniq.append(y)
        gaps = [b.yc - a.yc for a, b in zip(cd, cd[1:])]
        gap = statistics.median(gaps) if gaps else 70
        prev_bot = None
        for i, d in enumerate(cd):
            above = [y for y in uniq if y <= d.top + 2 and d.top - y < gap]
            below = [y for y in uniq if y > d.top + 2 and y - d.top < gap * 1.2]
            lead = gap * 0.28
            if above:
                top = above[-1] - 1
            else:  # 罫線が取れない場合の目安
                top = max(d.top - lead, prev_bot) if prev_bot else d.top - lead
            if below:
                bot = below[0] - 1
            else:
                bot = cd[i + 1].top - lead if i + 1 < len(cd) else min(d.top + gap, grid_bottom)
            cell_bounds[id(d)] = (max(top, grid_top), bot)
            prev_bot = bot

    order_x = {}
    for c in range(len(col_lefts)):
        ko = [s.x0 for s in spans if s.text.strip() == "個" and col_of(s.x0) == c and grid_top < s.y0 < grid_bottom]
        order_x[c] = (min(ko) - 12) if ko else col_lefts[c] + (col_right(c) - col_lefts[c]) * 0.66

    skip_words = {"個", "備考", "ヘルシー", "デラックス", "贅沢弁当"}
    for d in days:
        top, bot = cell_bounds[id(d)]
        cl, cr = col_lefts[d.col], order_x[d.col]
        cell = [s for s in spans if id(s) not in used and top <= s.yc < bot and cl - 6 <= s.x0 < cr
                and s.text.strip() not in skip_words and not (s.din and s.text.strip().isdigit() and s.size >= 15)]
        # アイコン（タブ）
        tab = next((t for t in tabs if col_of((t["rect"].x0 + t["rect"].x1) / 2) == d.col and top <= (t["rect"].y0 + t["rect"].y1) / 2 < bot), None)
        whites = [s for s in cell if s.white]
        type_txt = [s for s in whites if s.text.strip() in TYPE_WORDS
                    or (tab and tab["rect"].contains(pymupdf.Rect(s.x0, s.y0, s.x1, s.y1)))]
        if type_txt:
            d.icon_type = TYPE_ALIASES.get(type_txt[0].text.strip(), type_txt[0].text.strip())
            d.icon_how = "文字"
        elif tab:
            d.icon_type = _icon_from_glyphs(tab["glyphs"]) or "（判定不可）"
            d.icon_how = "図形"
        elif whites:
            d.icon_type = "（判定不可）"
        d.icon_name = "".join(s.text.strip() for s in sorted(whites, key=lambda s: (s.yc, s.x0)) if s not in type_txt)
        if d.icon_name and not d.icon_type:
            d.icon_type = "（判定不可）"
        cell = [s for s in cell if not s.white]
        # ご飯マーク
        marks = [s for s in cell if s.text.strip() in cfg.rice_words and s.size < 8.5]
        d.rice_marks = [m.text.strip() for m in marks]
        cell = [s for s in cell if s not in marks]
        # 贅沢弁当ラベル
        zlabel = [s for s in cell if s.text.strip() in ("贅沢", "弁当") and s.size <= 8]
        zx = max((s.x1 for s in zlabel), default=None)
        cell = [s for s in cell if s not in zlabel]
        lines = _group_lines(cell)
        has_bullet = any(s.text.strip().startswith("●") for s in cell)
        side_lines, z_lines = [], []
        cur_item = None
        for ln in lines:
            sp = ln["spans"]
            bullet = next((s for s in sp if s.text.strip().startswith("●")), None)
            rice_here = any(abs(m.yc - ln["yc"]) <= 3 for m in marks)
            if bullet:
                txt = _join([s for s in sp if s.x0 >= bullet.x0]).replace("●", "", 1)
                cur_item = [fix_extract(txt), False, rice_here]
                d.items.append(cur_item)
            elif all(s.regular for s in sp):
                side_lines.append(fix_extract(_join(sp)))
                cur_item = None
            elif zx is not None and sp[0].x0 > zx:
                z_lines.append((_join(sp), rice_here))
            elif cur_item is not None and has_bullet:
                cur_item[0] = fix_extract(cur_item[0] + _join(sp))
                cur_item[1] = True
            else:  # 「●」が付いていないレイアウト：1行＝1品
                cur_item = [fix_extract(_join(sp)), False, rice_here]
                d.items.append(cur_item)
        d.side_lines = side_lines
        raw = "".join(side_lines)
        d.sides_raw = raw
        d.sides = [x.strip() for x in raw.split("/")] if raw else []
        if z_lines:
            joined = "\n".join(fix_extract(t) for t, _ in z_lines)
            d.zeitaku_has_slash = "/" in joined
            parts = joined.split("/") if d.zeitaku_has_slash else [joined]
            d.zeitaku = [(p.replace("\n", "").strip(), "\n" in p.strip()) for p in parts if p.strip()]
    days.sort(key=lambda d: (d.year, d.month, d.day))

    # --- イチオシ ---
    ichioshi = []
    entries = []
    for s in din_digits:
        if s.y0 >= grid_top or id(s) in used:
            continue
        wd = [w for w in weekday_spans if w.y0 < grid_top and s.x1 - 4 <= w.x0 <= s.x1 + 12
              and w.y0 < s.y1 and w.y1 > s.y0]
        if wd:
            entries.append((s, wd[0]))
    entries.sort(key=lambda e: e[0].yc)
    split_x = col_lefts[1] if len(col_lefts) > 1 else page_w * 0.5
    type_all = [x for x in spans if x.text.strip() in TYPE_WORDS and x.y0 < grid_top]
    tops = []
    for s, w in entries:
        tl = [x for x in type_all if s.y0 - 35 <= x.y0 <= s.y0 + 5 and x.x0 < s.x1 + 10]
        tops.append(min(x.y0 for x in tl) - 8 if tl else None)
    for i, (s, w) in enumerate(entries):
        if tops[i] is not None:
            top = tops[i]
        elif i == 0:
            top = s.y0 - 45
        else:
            top = (entries[i - 1][0].y0 + s.y0) / 2 - 5
        if i + 1 < len(entries):
            nxt = entries[i + 1][0]
            bot = tops[i + 1] if tops[i + 1] is not None else (s.y0 + nxt.y0) / 2 - 5
        else:
            bot = min(s.y1 + 45, grid_top)
        reg = [x for x in spans if top <= x.yc < bot and x.x0 < split_x and x is not s and x is not w and x.y0 < grid_top]
        pref = [p for p in reg if p.din and p.text.strip().isdigit() and p.x1 <= s.x0 + 1 and s.x0 - p.x1 < 25
                and abs(p.yc - s.yc) < s.size]
        month = int(pref[0].text.strip()) if pref else None
        reg = [x for x in reg if not (x.din and x.text.strip().isdigit())]
        type_sp = [x for x in reg if x.text.strip() in TYPE_WORDS]
        itype = TYPE_ALIASES.get(type_sp[0].text.strip(), "") if type_sp else ""
        reg = [x for x in reg if x not in type_sp and not x.white and x.size < 20
               and not any(h in x.text for h in ("おすすめメニュー", "イチオシ", "今月の"))]
        comments = [x for x in reg if x.regular and x.x0 > w.x1 and x.text.strip() != "●"]
        menu_x = min((x.x0 for x in comments), default=None)
        bullets = [x for x in reg if x.text.strip() == "●" and x.x0 > w.x1
                   and (menu_x is None or abs(x.x0 - menu_x) < 15)]
        if bullets:
            menu_x = min(menu_x or 1e9, min(b.x0 for b in bullets))
        if menu_x is None:
            continue
        body = [x for x in reg if x.x0 >= menu_x - 8 and x.text.strip() != "●"]
        blines = [l for l in _group_lines(body) if l["spans"][0].x0 <= menu_x + 40]
        # 日付に一番近い行から、行間が詰まっている範囲だけを採用（欄外の文字を除く）
        if blines:
            k = min(range(len(blines)), key=lambda i: abs(blines[i]["yc"] - s.yc))
            lo = hi = k

            def _h(l):
                return max(x.y1 for x in l["spans"]) - min(x.y0 for x in l["spans"])

            while lo > 0 and min(x.y0 for x in blines[lo]["spans"]) - max(x.y1 for x in blines[lo - 1]["spans"]) < max(_h(blines[lo]), 8) * 0.9:
                lo -= 1
            while hi + 1 < len(blines) and min(x.y0 for x in blines[hi + 1]["spans"]) - max(x.y1 for x in blines[hi]["spans"]) < max(_h(blines[hi]), 8) * 0.9:
                hi += 1
            blines = blines[lo:hi + 1]
        mlines = [l for l in blines if not all(x.regular for x in l["spans"])]
        clines = [l for l in blines if all(x.regular for x in l["spans"])]
        event_sp = [x for x in reg if not x.regular and w.x1 - 2 <= x.x0 < menu_x - 8
                    and s.y0 - 15 <= x.y0 <= s.y1 + 5]
        menu = fix_extract("".join(_join(l["spans"]).strip() for l in mlines))
        comment = fix_extract("".join(_join(l["spans"]) for l in clines))
        event = fix_extract("".join(_join(l["spans"]) for l in _group_lines(event_sp)))
        ichioshi.append(Ichioshi(int(s.text.strip()), month, w.text.strip(), itype, event, menu,
                                 len(mlines) > 1, comment))
    for it in ichioshi:
        cands = [d for d in days if d.day == it.day and (it.month is None or d.month == it.month)]
        if cands:
            it.month, it.year = cands[0].month, cands[0].year
        elif it.month:
            it.year = days[0].year if it.month >= days[0].month else days[0].year + 1

    footnote_text = "".join(fix_extract(s.text) for s in spans if s.y0 >= grid_bottom - 2)
    return PdfDoc(name, title, zeitaku_edition, period, days, ichioshi, fulltext, footnote_text, warnings)


# ----------------------------------------------------------------------------
# チェック
# ----------------------------------------------------------------------------
@dataclass
class Finding:
    severity: str
    category: str
    date: str
    weekday: str
    item: str
    pdf: str
    master: str
    note: str


def fmt_date(d: Optional[dt.date]) -> str:
    return f"{d.month}/{d.day}" if d else "-"


def run_checks(pdf: PdfDoc, masters: dict, fax: list, cfg: Config):
    F = []
    status = {}

    def add(sev, cat, d, wd, item, p, m, note=""):
        F.append(Finding(sev, cat, fmt_date(d), wd, item, p, m, note))

    def mark(cat, ok, detail=""):
        prev = status.get(cat)
        if prev is None or (prev[0] == "OK" and not ok):
            status[cat] = ("OK" if ok else "要確認", detail)

    CAT1, CAT2, CAT3, CAT4 = "①曜日", "②メニュー名", "③アイコン", "④イチオシ"
    CAT5, CAT6, CAT7, CAT8 = "⑤贅沢あり：定型文", "⑥贅沢あり：混ぜご飯・混ぜ麺・ふりかけ", "⑦掲載日・期間", "⑧共通の定型文言"
    for c in [CAT1, CAT2, CAT3, CAT4, CAT7, CAT8]:
        status[c] = ("OK", "")
    if pdf.zeitaku:
        status[CAT5] = ("OK", "")
        status[CAT6] = ("OK", "")

    # ① 曜日
    for d in pdf.days:
        if d.date is None:
            add("高", CAT1, None, d.weekday_char, "日付", f"{d.month}/{d.day}", "", "存在しない日付です")
            mark(CAT1, False)
            continue
        exp = d.date.weekday()
        got = WEEKDAY_CHARS.get(d.weekday_char)
        if got != exp:
            add("高", CAT1, d.date, d.weekday_char, "曜日", d.weekday_char, WEEKDAY_JA[exp], "日付と曜日が合っていません")
            mark(CAT1, False)

    # ② メニュー名 ③ アイコン ⑥ ご飯
    used_rice_words = set()
    for d in pdf.days:
        md: Optional[MasterDay] = masters.get(d.date)
        wd = d.weekday_char
        if md is None:
            add("高", CAT7, d.date, wd, "日付", f"{fmt_date(d.date)}", "", "マスタにこの日付がありません")
            mark(CAT7, False)
            continue
        # メニュー
        items = list(d.items)
        extra = [it for it in items if it[2]]  # ご飯マーク付きの行（7/7 鮭ちらし寿司 など）
        items = [it for it in items if not it[2]]
        pdf_vals = {}
        if len(items) >= 1:
            pdf_vals["main"] = (items[0][0], items[0][1])
        if len(items) >= 2:
            pdf_vals["submain"] = (items[1][0], items[1][1])
        for it in items[2:]:
            add("中", CAT2, d.date, wd, "メニュー行", it[0], "", "メイン・サブメイン以外の「●」行があります")
            mark(CAT2, False)
        if not pdf.zeitaku:
            mvals = [md.fields.get(f, "") for f in ["fukusai1", "fukusai2", "salad"]]
            sides, by_break = split_sides(d.side_lines, mvals) if d.side_lines else ([], False)
            for f, v in zip(["fukusai1", "fukusai2", "salad"], sides):
                pdf_vals[f] = v
            if by_break:
                d.sides = [x[0] for x in sides]
            if by_break:
                add("低", CAT2, d.date, wd, "副菜・サラダ", " / ".join(d.side_lines), " / ".join(mvals),
                    "区切りの「/」がなく改行で区切られている箇所があります")
                mark(CAT2, False)
            elif sides and len(sides) != 3:
                add("中", CAT2, d.date, wd, "副菜・サラダ", d.sides_raw, " / ".join(mvals),
                    f"「/」区切りが{len(sides)}項目です（3項目のはず）")
                mark(CAT2, False)
            fields = MENU_FIELDS
        else:
            for f, (v, wr) in zip(ZEITAKU_FIELDS, d.zeitaku):
                pdf_vals[f] = (v, wr)
            # 「/」がなく、2行がそれぞれ贅沢1・2に当たる場合
            if not d.zeitaku_has_slash and len(d.zeitaku) == 1 and md.fields.get("zeitaku2", "").strip():
                joined = d.zeitaku[0][0]
                z1, z2 = md.fields["zeitaku1"], md.fields["zeitaku2"]
                if no_ws(joined) == no_ws(z1 + z2):
                    pdf_vals["zeitaku1"] = (z1, True)
                    pdf_vals["zeitaku2"] = (z2, True)
                    add("低", CAT2, d.date, wd, "贅沢弁当", joined, f"{z1} / {z2}",
                        "贅沢弁当1と2の間に区切りの「/」がありません（他の日は「/」区切り）")
                    mark(CAT2, False)
            fields = ["main", "submain"] + ZEITAKU_FIELDS
        for f in fields:
            mv = md.fields.get(f, "")
            pv = pdf_vals.get(f)
            if pv is None:
                if mv.strip():
                    add("高", CAT2, d.date, wd, FIELD_JA[f], "（記載なし）", mv, "PDFに見当たりません")
                    mark(CAT2, False)
                continue
            if not mv.strip():
                add("高", CAT2, d.date, wd, FIELD_JA[f], pv[0], "（マスタは空欄）", "マスタにない項目です")
                mark(CAT2, False)
                continue
            res, note = compare_text(pv[0], mv, pv[1])
            if res != "OK":
                add(res, CAT2, d.date, wd, FIELD_JA[f], describe_ws(fix_extract(pv[0])) if res == "低" else fix_extract(pv[0]),
                    describe_ws(mv.strip()) if res == "低" else mv.strip(), note)
                mark(CAT2, False)
        for it in extra:
            row_texts = [md.memo] + list(md.fields.values())
            hit = any(it[0] and no_ws(it[0]) in no_ws(t) for t in row_texts if t)
            add("情報", CAT2, d.date, wd, "ご飯マーク付きの行", it[0],
                md.memo or "", "マスタの行内（メモ列など）に記載あり" if hit else "マスタの行内に見当たりません")
            if not hit:
                mark(CAT2, False)

        # アイコン
        exp_type = md.icon_type
        exp_note = ""
        if exp_type == "贅沢プレミアム" and not pdf.zeitaku:
            exp_type, exp_note = None, "（贅沢プレミアムは贅沢なし版では表示しない扱い）"
        got_type = d.icon_type
        if (exp_type or None) != (got_type or None):
            add("高", CAT3, d.date, wd, "アイコン種別", got_type or "（なし）",
                f"{exp_type or '（なし）'}" + (f"　※A列:{md.label}" if md.label else ""),
                ("マスタの色: #" + md.label_color if md.label_color else "") + exp_note)
            mark(CAT3, False)
        elif exp_type and exp_type in ("プレミアム", "フェア") and md.event_name:
            res, note = compare_text(d.icon_name, md.event_name)
            if res != "OK":
                fx = next((r for r in fax if r.month == d.month and r.day == d.day), None)
                extra_note = ""
                if fx and fx.event and no_ws(fx.event) == no_ws(d.icon_name):
                    extra_note = "（FaxDM訴求文のイベント名とは一致）"
                add("中" if res == "高" else res, CAT3, d.date, wd, "アイコン名", d.icon_name or "（なし）",
                    md.event_name, (note or "アイコン横の名前がマスタと違います") + extra_note)
                mark(CAT3, False)

        # ご飯
        if pdf.zeitaku:
            if md.rice_word:
                used_rice_words.add(md.rice_word)
                if md.rice_word not in d.rice_marks:
                    add("高", CAT6, d.date, wd, "ご飯マーク", ", ".join(d.rice_marks) or "（なし）", md.rice_word,
                        "マスタでは" + md.label + "の日です")
                    mark(CAT6, False)
            for mk in d.rice_marks:
                if mk != md.rice_word and not (md.rice_word and md.rice_word not in d.rice_marks):
                    add("中", CAT6, d.date, wd, "ご飯マーク", mk, md.rice_word or "（なし）",
                        "マスタでは" + (md.label or "該当しない日") + "です")
                    mark(CAT6, False)
                used_rice_words.add(mk)
        elif d.rice_marks:
            add("情報", CAT6 if pdf.zeitaku else CAT2, d.date, wd, "ご飯マーク", ", ".join(d.rice_marks),
                md.rice_word or md.memo or "", "贅沢なし版にご飯マークがあります（必要なら確認）")

    # ⑥ 注記
    if pdf.zeitaku:
        for w in sorted(used_rice_words):
            phrase = f"{w}の日は、ご飯の変更はできません。"
            if no_ws(phrase) not in no_ws(pdf.fulltext):
                add("高", CAT6, None, "", "注記", "（なし）", phrase, "欄外の注記が見当たりません")
                mark(CAT6, False)
        if not used_rice_words:
            status[CAT6] = ("OK", "掲載期間に該当日なし")

    # ⑤ 贅沢あり定型文
    if pdf.zeitaku:
        for ph in cfg.zeitaku_required:
            if no_ws(ph) not in no_ws(pdf.fulltext):
                add("高", CAT5, None, "", "定型文", "（なし）", ph, "贅沢あり版に必要な文言が見当たりません")
                mark(CAT5, False)
    # ⑧ 共通文言
    for ph in cfg.common_required:
        target = no_ws(ph)
        text = no_ws(pdf.fulltext)
        if target in text:
            continue
        head = target[:3]
        m = re.search(re.escape(head) + ".{0," + str(len(target)) + "}", text)
        if m:
            add("高", CAT8, None, "", "定型文言", m.group(0)[:len(target)], ph, "定型文言が違います（誤字の可能性）")
            mark(CAT8, False)
        # 文字として読み取れない（画像・アウトライン化）場合は何も出さない

    # ⑦ 期間・掲載漏れ
    if pdf.days:
        first, last = pdf.days[0].date, pdf.days[-1].date
        if pdf.period:
            _, ps, pe = pdf.period
            if first != ps:
                add("中", CAT7, first, "", "掲載開始日", fmt_date(first), fmt_date(ps), "ファイル名の掲載期間と最初の日付が違います")
                mark(CAT7, False)
            if last != pe:
                add("中", CAT7, last, "", "掲載終了日", fmt_date(last), fmt_date(pe), "ファイル名の掲載期間と最後の日付が違います")
                mark(CAT7, False)
        pdf_dates = {d.date for d in pdf.days}
        for dd, md in sorted(masters.items()):
            if first <= dd <= last and md.is_full and dd not in pdf_dates:
                add("高", CAT7, dd, WEEKDAY_JA[dd.weekday()], "掲載漏れ", "（なし）", md.fields.get("main", ""),
                    "マスタにある日がPDFにありません")
                mark(CAT7, False)
        seq = [(d.col, d.yc, d.date) for d in sorted(pdf.days, key=lambda d: (d.col, d.yc))]
        for a, b in zip(seq, seq[1:]):
            if a[2] and b[2] and b[2] <= a[2]:
                add("中", CAT7, b[2], "", "並び順", fmt_date(b[2]), "", f"{fmt_date(a[2])} の次に並んでいます")
                mark(CAT7, False)

    # ④ イチオシ
    if not pdf.ichioshi:
        status[CAT4] = ("OK", "イチオシ欄なし")
    for it in pdf.ichioshi:
        d = it.date
        wd = it.weekday_char
        if d is None:
            add("高", CAT4, None, wd, "日付", f"{it.month}/{it.day}", "", "注文欄に同じ日付がなく、月を特定できません")
            mark(CAT4, False)
            continue
        if WEEKDAY_CHARS.get(wd) != d.weekday():
            add("高", CAT4, d, wd, "曜日", wd, WEEKDAY_JA[d.weekday()], "イチオシ欄の曜日が日付と合っていません")
            mark(CAT4, False)
        gd = next((x for x in pdf.days if x.date == d), None)
        md = masters.get(d)
        if gd and gd.items:
            gmain = next((x[0] for x in gd.items if not x[2]), "")
            res, _ = compare_text(it.menu, gmain, True)
            if res == "高":
                add("高", CAT4, d, wd, "メニュー（注文欄と比較）", it.menu, gmain, "イチオシ欄と注文欄のメインが違います")
                mark(CAT4, False)
        if md:
            res, _ = compare_text(it.menu, md.fields.get("main", ""), True)
            if res != "OK":
                add(res, CAT4, d, wd, "メニュー（マスタと比較）", it.menu, md.fields.get("main", ""), "")
                mark(CAT4, False)
            exp_t = md.icon_type
            if it.type and exp_t and it.type != exp_t:
                add("高", CAT4, d, wd, "種別", it.type, exp_t, "マスタの色と違います")
                mark(CAT4, False)
            if gd and it.type and gd.icon_type and it.type != gd.icon_type:
                add("中", CAT4, d, wd, "種別（注文欄と比較）", it.type, gd.icon_type, "")
                mark(CAT4, False)
        fx = [r for r in fax if r.month == d.month and r.day == d.day]
        if not fax:
            continue
        if not fx:
            add("高", CAT4, d, wd, "訴求文", it.comment, "（FaxDM訴求文に行なし）", "")
            mark(CAT4, False)
            continue
        r = next((r for r in fx if r.fixed), fx[0])
        if not r.fixed:
            add("高", CAT4, d, wd, "一言コメント", it.comment, "（確定が空欄）", "FaxDM訴求文の「確定」が未入力です")
            mark(CAT4, False)
        else:
            res, note = compare_text(it.comment, r.fixed, True)
            if res != "OK":
                add(res, CAT4, d, wd, "一言コメント", it.comment, r.fixed, note)
                mark(CAT4, False)
        res, _ = compare_text(it.menu, r.menu, True)
        if res != "OK":
            add("中" if res == "高" else res, CAT4, d, wd, "メニュー（FaxDM訴求文と比較）", it.menu, r.menu, "")
            mark(CAT4, False)
        if it.type and r.type and it.type != r.type:
            add("中", CAT4, d, wd, "種別（FaxDM訴求文と比較）", it.type, r.type, "")
            mark(CAT4, False)
        if it.event and r.event and compare_text(it.event, r.event, True)[0] != "OK":
            add("中", CAT4, d, wd, "イベント名（FaxDM訴求文と比較）", it.event, r.event, "")
            mark(CAT4, False)
        if r.adopt == "否":
            add("中", CAT4, d, wd, "選定可否", "イチオシに掲載", "否", "FaxDM訴求文では「否」になっています")
            mark(CAT4, False)
    F.sort(key=lambda f: (SEV_ORDER.get(f.severity, 9), f.category, f.date))
    return F, status


# ----------------------------------------------------------------------------
# Excel 出力
# ----------------------------------------------------------------------------
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
SEV_FILL = {"高": "F8CBAD", "中": "FFE699", "低": "E2EFDA", "情報": "DDEBF7"}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "Arial"


def _table(ws, headers, rows, widths, start_row=1):
    for j, h in enumerate(headers, 1):
        c = ws.cell(start_row, j, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for i, r in enumerate(rows, start_row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name=FONT)
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{ws.cell(start_row + len(rows), len(headers)).coordinate}"


def write_report(pdf: PdfDoc, findings: list, status: dict, masters: dict, notes: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "サマリー"
    ws["A1"] = "FAXDM 校正結果"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    info = [
        ("ファイル", pdf.path),
        ("版", "贅沢あり" if pdf.zeitaku else "贅沢なし"),
        ("掲載日", f"{fmt_date(pdf.days[0].date)}〜{fmt_date(pdf.days[-1].date)}（{len(pdf.days)}日）" if pdf.days else "-"),
        ("ファイル名の掲載期間", f"{fmt_date(pdf.period[1])}〜{fmt_date(pdf.period[2])}" if pdf.period else "読み取れず"),
        ("参照マスタ", ", ".join(sorted({m.source for m in masters.values()})) or "なし"),
        ("実行日時", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(i, 1, k).font = Font(name=FONT, bold=True)
        ws.cell(i, 2, v).font = Font(name=FONT)
    r0 = len(info) + 4
    counts = {s: sum(1 for f in findings if f.severity == s) for s in SEV_ORDER}
    ws.cell(r0, 1, "指摘件数").font = Font(name=FONT, bold=True)
    ws.cell(r0, 2, "　".join(f"{k}:{v}" for k, v in counts.items())).font = Font(name=FONT)
    rows = []
    for cat, (st, det) in sorted(status.items()):
        n = sum(1 for f in findings if f.category == cat and f.severity != "情報")
        rows.append([cat, st, det or (f"{n}件" if n else "")])
    _table(ws, ["チェック項目", "結果", "詳細"], rows, [34, 12, 60], start_row=r0 + 2)
    for i in range(r0 + 3, r0 + 3 + len(rows)):
        c = ws.cell(i, 2)
        c.fill = PatternFill("solid", fgColor="E2EFDA" if c.value == "OK" else "F8CBAD")
        c.alignment = Alignment(horizontal="center", vertical="center")
    rr = r0 + 4 + len(rows)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 50
    for msg in notes + pdf.warnings:
        ws.cell(rr, 1, "※ " + msg).font = Font(name=FONT, color="C00000")
        rr += 1
    ws.cell(rr + 1, 1, "重要度　高＝誤りの可能性大／中＝要確認（運用どおりなら問題なし）／低＝スペースや記号の細部／情報＝参考").font = \
        Font(name=FONT, italic=True, color="595959")
    ws.cell(rr + 2, 1, "「1 品」「4 種」のような数字の後ろの空白は、PDF抽出時に入るものとして無視しています。").font = \
        Font(name=FONT, italic=True, color="595959")

    # 指摘一覧
    ws2 = wb.create_sheet("指摘一覧")
    rows = [[i, f.severity, f.category, f.date, f.weekday, f.item, f.pdf, f.master, f.note, "未対応"]
            for i, f in enumerate(findings, 1)]
    _table(ws2, ["No", "重要度", "チェック項目", "日付", "曜日", "欄", "PDFの表記", "マスタの表記", "備考", "対応状況"],
           rows, [5, 8, 22, 8, 6, 18, 36, 36, 40, 12])
    for i, f in enumerate(findings, 2):
        ws2.cell(i, 2).fill = PatternFill("solid", fgColor=SEV_FILL.get(f.severity, "FFFFFF"))
        for c in (1, 2, 4, 5, 10):
            ws2.cell(i, c).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(i, 10).fill = PatternFill("solid", fgColor="FFFF00")
    if findings:
        dv = DataValidation(type="list", formula1='"未対応,修正済,問題なし（運用どおり）"', allow_blank=True)
        ws2.add_data_validation(dv)
        dv.add(f"J2:J{len(findings) + 1}")
    else:
        ws2.cell(2, 1, "指摘はありません").font = Font(name=FONT)

    # 読み取り結果
    ws3 = wb.create_sheet("PDF読み取り結果")
    rows = []
    for d in pdf.days:
        items = [it for it in d.items if not it[2]]
        extra = " / ".join(it[0] for it in d.items if it[2])
        row = [fmt_date(d.date), d.weekday_char, d.icon_type or "", d.icon_name,
               items[0][0] if items else "", items[1][0] if len(items) > 1 else ""]
        if pdf.zeitaku:
            row += [d.zeitaku[0][0] if d.zeitaku else "", d.zeitaku[1][0] if len(d.zeitaku) > 1 else ""]
        else:
            row += [(d.sides + ["", "", ""])[k] for k in range(3)]
        row += [", ".join(d.rice_marks), extra]
        rows.append(row)
    hdr = ["日付", "曜日", "アイコン", "アイコン名", "メイン", "サブメイン"] + \
          (["贅沢弁当1", "贅沢弁当2"] if pdf.zeitaku else ["副菜1", "副菜2", "サラダ(デザート)"]) + ["ご飯マーク", "マーク付きの行"]
    _table(ws3, hdr, rows, [8, 6, 12, 18, 30, 30] + ([34, 34] if pdf.zeitaku else [26, 26, 28]) + [12, 20])

    ws4 = wb.create_sheet("イチオシ")
    rows = [[fmt_date(it.date), it.weekday_char, it.type, it.event, it.menu, it.comment] for it in pdf.ichioshi]
    _table(ws4, ["日付", "曜日", "種別", "イベント名", "メニュー", "一言コメント"], rows, [8, 6, 12, 22, 36, 60])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# 校正の実行
# ----------------------------------------------------------------------------
def run_one(name: str, data: bytes, src: MasterSource, cfg: Config) -> dict:
    src.notes = []
    pdf = read_pdf(data, name, cfg)
    masters = {}
    for y, m in sorted({(d.year, d.month) for d in pdf.days if d.year}):
        masters.update(src.monthly(y, m))
    fax = src.faxdm()
    findings, status = run_checks(pdf, masters, fax, cfg)
    notes = list(dict.fromkeys(src.notes))
    excel = write_report(pdf, findings, status, masters, notes)
    return {"name": name, "pdf": pdf, "findings": findings, "status": status, "notes": notes,
            "masters": sorted({m.source for m in masters.values()}), "excel": excel, "error": None}


# ----------------------------------------------------------------------------
# 画面
# ----------------------------------------------------------------------------
SEV_ICON = {"高": "🔴 高", "中": "🟠 中", "低": "🟢 低", "情報": "🔵 情報"}


def _secrets() -> dict:
    try:
        return dict(st.secrets)
    except Exception:  # secrets.toml がない
        return {}


def show_result(r: dict):
    if r["error"]:
        st.error(f"読み取れませんでした：{r['error']}")
        return
    pdf, findings = r["pdf"], r["findings"]
    days = pdf.days
    st.caption(
        f"{'贅沢あり' if pdf.zeitaku else '贅沢なし'}版　｜　掲載日 {fmt_date(days[0].date)}〜{fmt_date(days[-1].date)}（{len(days)}日）"
        + (f"　｜　ファイル名の期間 {fmt_date(pdf.period[1])}〜{fmt_date(pdf.period[2])}" if pdf.period else "")
        + f"　｜　参照マスタ：{', '.join(r['masters']) or 'なし'}")
    for n in r["notes"] + pdf.warnings:
        st.warning(n)

    cols = st.columns(4)
    for c, sev in zip(cols, ["高", "中", "低", "情報"]):
        c.metric(SEV_ICON[sev], sum(1 for f in findings if f.severity == sev))

    st.markdown("**指摘一覧**")
    sel = st.multiselect("表示する重要度", ["高", "中", "低", "情報"], default=["高", "中", "低"],
                         key=f"sev_{r['name']}")
    rows = [{"重要度": SEV_ICON[f.severity], "項目": f.category, "日付": f.date, "曜日": f.weekday,
             "欄": f.item, "PDFの表記": f.pdf, "マスタの表記": f.master, "備考": f.note}
            for f in findings if f.severity in sel]
    if rows:
        st.dataframe(pd.DataFrame(rows), hide_index=True, use_container_width=True)
    else:
        st.success("指摘はありません")

    with st.expander("チェック項目ごとの結果", expanded=True):
        st.dataframe(pd.DataFrame([{"項目": k, "結果": ("✅ OK" if v[0] == "OK" else "⚠️ 要確認"),
                                    "詳細": v[1]} for k, v in sorted(r["status"].items())]),
                     hide_index=True, use_container_width=True)

    with st.expander("PDFから読み取った内容（読み取りが正しいかの確認用）"):
        rows = []
        for d in days:
            items = [it for it in d.items if not it[2]]
            row = {"日付": fmt_date(d.date), "曜日": d.weekday_char, "アイコン": d.icon_type or "",
                   "アイコン名": d.icon_name, "メイン": items[0][0] if items else "",
                   "サブメイン": items[1][0] if len(items) > 1 else ""}
            if pdf.zeitaku:
                row["贅沢弁当1"] = d.zeitaku[0][0] if d.zeitaku else ""
                row["贅沢弁当2"] = d.zeitaku[1][0] if len(d.zeitaku) > 1 else ""
            else:
                for k, lab in enumerate(["副菜1", "副菜2", "サラダ(デザート)"]):
                    row[lab] = d.sides[k] if k < len(d.sides) else ""
            row["ご飯マーク"] = ", ".join(d.rice_marks)
            rows.append(row)
        st.dataframe(pd.DataFrame(rows), hide_index=True, use_container_width=True)
        if pdf.ichioshi:
            st.markdown("**イチオシ**")
            st.dataframe(pd.DataFrame([{"日付": fmt_date(i.date), "曜日": i.weekday_char, "種別": i.type,
                                        "イベント名": i.event, "メニュー": i.menu, "一言コメント": i.comment}
                                       for i in pdf.ichioshi]), hide_index=True, use_container_width=True)

    base = r["name"].rsplit(".", 1)[0]
    st.download_button("📥 結果をExcelでダウンロード", r["excel"], file_name=f"校正結果_{base}.xlsx",
                       mime=XLSX_MIME, key=f"dl_{r['name']}")


def main():
    st.set_page_config(page_title="FAXDM 校正ツール", page_icon="🍱", layout="wide")
    st.title("🍱 FAXDM 校正ツール")
    st.caption("PDF原稿をマスタ（YYMM_やどかり／FaxDM訴求文）と照合します。マスタは各ファイルの一番左のシートを参照します。")

    sec = _secrets()
    has_sa = "gcp_service_account" in sec
    cfg = Config()

    with st.sidebar:
        st.header("マスタ")
        modes = ["スプレッドシートから読み込む", "ファイルをアップロード"]
        if has_sa:
            modes.insert(0, "Googleドライブから自動取得")
        mode = st.radio("取得方法", modes)
        if mode == "スプレッドシートから読み込む":
            months = sorted(MASTER_SHEETS)
            st.caption("PDFの掲載期間に合わせて、必要な月のマスタと FaxDM訴求文 を自動で読み込みます"
                       "（月をまたぐ場合は両方の月）。")
            st.caption("登録済み：" + (f"{months[0][:4]}〜{months[-1][:4]}" if months else "なし")
                       + "　／　FaxDM訴求文")
            extra = st.text_area("未登録の月のURL（1行に1つ）", height=90,
                                 placeholder="https://docs.google.com/spreadsheets/d/.../edit",
                                 help="MASTER_SHEETS に登録していない月だけ貼ってください。"
                                      "毎月使うなら app.py の MASTER_SHEETS に1行追加すると、貼らなくて済みます。")
            faxdm_url = sec.get("faxdm_sheet_id", FAXDM_SHEET_ID)
            provider = sheet_provider(MASTER_SHEETS, extra.splitlines(), faxdm_url, cfg.faxdm_name)
            st.caption("※ シートの共有設定は「リンクを知っている全員」が「閲覧者」にしてください。")
            if st.button("🔄 マスタを今すぐ取り直す"):
                st.cache_data.clear()
                st.toast("次の校正で最新のマスタを取得します")
        elif mode == "ファイルをアップロード":
            files = st.file_uploader(
                "マスタ（.xlsx）", type=["xlsx"], accept_multiple_files=True,
                help="Googleスプレッドシートを「ファイル → ダウンロード → Microsoft Excel」で保存したもの。"
                     "例：2610_やどかり.xlsx、2611_やどかり.xlsx、FaxDM訴求文.xlsx")
            provider = upload_provider({f.name: f.getvalue() for f in files or []})
        else:
            folder = sec.get("drive_folder_id", DRIVE_FOLDER_ID)
            faxdm_id = sec.get("faxdm_sheet_id", FAXDM_SHEET_ID)
            provider = drive_provider(sec["gcp_service_account"], folder, faxdm_id, cfg.faxdm_name)
            st.caption("最新のマスタを取得します（10分間は同じ内容を使い回します）。")
            if st.button("🔄 マスタを今すぐ取り直す"):
                st.cache_data.clear()
                st.toast("次の校正で最新のマスタを取得します")
        with st.expander("判定ルール"):
            st.markdown("**アイコンの色（ラベル列）**")
            for k, v in ICON_COLORS.items():
                st.markdown(f"<span style='background:#{k};padding:0 10px;border:1px solid #ccc'>&nbsp;</span> "
                            f"#{k} → {v}", unsafe_allow_html=True)
            st.markdown("**ご飯変更不可の日**：ラベルに " + "・".join(RICE_WORDS) + " を含む日")
            st.markdown("**重要度**：高＝誤りの可能性大／中＝要確認／低＝スペースや記号の細部／情報＝参考")

    pdfs = st.file_uploader("PDF原稿（複数可）", type=["pdf"], accept_multiple_files=True)
    if st.button("校正する", type="primary", disabled=not pdfs):
        src = MasterSource(provider, cfg)
        results = []
        with st.spinner("校正中…"):
            for f in pdfs:
                try:
                    results.append(run_one(f.name, f.getvalue(), src, cfg))
                except Exception as e:
                    results.append({"name": f.name, "error": str(e)})
        st.session_state["results"] = results

    results = st.session_state.get("results")
    if results:
        tabs = st.tabs([r["name"] for r in results]) if len(results) > 1 else [st.container()]
        for t, r in zip(tabs, results):
            with t:
                if len(results) == 1:
                    st.subheader(r["name"])
                show_result(r)


main()
